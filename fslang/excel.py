"""讀取多語系翻譯 Excel（以 Excel 為標準答案）。

預期格式（與 FS - Version 137 相同）：
- 第 1 列為表頭，每一欄一個語言，例如 EN / ITALY / Thai / TH (revised) / Japan …
- 之後每一列為一句文字，EN 欄為原文。
- 「XX (revised)」欄是前一個語言欄的修訂版：有填值時以修訂版為準，原欄文字視為「舊翻譯」。
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

from .config import GameConfig

# 表頭名稱（小寫）-> 語言代碼。設定檔的 "languages" 可覆寫。
DEFAULT_LANG_CODES = {
    "en": "en", "eng": "en", "english": "en",
    "italy": "it", "italian": "it", "it": "it",
    "thai": "th", "th": "th", "thailand": "th",
    "japan": "ja", "japanese": "ja", "jp": "ja", "ja": "ja",
    "korea": "ko", "korean": "ko", "kr": "ko", "ko": "ko",
    "turkish": "tr", "turkey": "tr", "tr": "tr",
    "russia": "ru", "russian": "ru", "ru": "ru",
    "spain": "es", "spanish": "es", "es": "es",
    "french": "fr", "france": "fr", "fr": "fr",
    "portugal": "pt", "portuguese": "pt", "pt": "pt",
    "brazil": "pt-br", "brazilian": "pt-br",
    "hindi": "hi", "india": "hi", "hi": "hi",
    "bangladesh": "bn", "bengali": "bn", "bangla": "bn", "bn": "bn",
    "german": "de", "germany": "de", "de": "de",
    "vietnam": "vi", "vietnamese": "vi", "vi": "vi",
    "indonesia": "id", "indonesian": "id", "id": "id",
    "malay": "ms", "malaysia": "ms", "ms": "ms",
    "myanmar": "my", "burmese": "my",
    "khmer": "km", "cambodia": "km",
    "chinese": "zh", "china": "zh", "sc": "zh-cn", "tc": "zh-tw",
    "簡中": "zh-cn", "簡體": "zh-cn", "繁中": "zh-tw", "繁體": "zh-tw",
}

_REVISED_RE = re.compile(r"\(?\s*revised\s*\)?|\(?\s*修正\s*\)?|\(?\s*修訂\s*\)?", re.I)


@dataclass
class Language:
    code: str            # 語言代碼，例如 th
    name: str            # Excel 表頭，例如 Thai
    column: int          # 欄號（1 起算）
    revision_column: int | None = None
    revision_name: str | None = None


@dataclass
class Translation:
    lang: str
    text: str | None         # 標準答案（有修訂版時為修訂版）
    cell: str                # 標準答案所在儲存格，例如 D3
    column_name: str         # 標準答案所在欄的表頭
    old_text: str | None = None   # 被修訂版取代的舊翻譯
    old_cell: str | None = None


@dataclass
class Entry:
    row: int
    key: str                 # 原文（EN）
    translations: dict[str, Translation] = field(default_factory=dict)


@dataclass
class TranslationSheet:
    path: Path
    sheet: str
    source_lang: str
    languages: list[Language]
    entries: list[Entry]

    def language(self, code: str) -> Language:
        for lang in self.languages:
            if lang.code == code:
                return lang
        raise KeyError(code)

    @property
    def codes(self) -> list[str]:
        return [lang.code for lang in self.languages]


def _clean(value) -> str | None:
    if value is None:
        return None
    text = str(value)
    return text if text.strip() else None


def _lookup(mapping: dict[str, str], header: str) -> str | None:
    key = header.strip().lower()
    for k, v in mapping.items():
        if k.strip().lower() == key:
            return v
    return None


def list_sheets(path: str | Path) -> list[str]:
    wb = openpyxl.load_workbook(path, read_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def load_sheet(cfg: GameConfig) -> TranslationSheet:
    wb = openpyxl.load_workbook(cfg.excel, data_only=True)
    try:
        if cfg.sheet:
            if cfg.sheet not in wb.sheetnames:
                raise ValueError(f"Excel 裡找不到工作表 '{cfg.sheet}'，現有：{wb.sheetnames}")
            ws = wb[cfg.sheet]
        else:
            ws = wb.worksheets[0]

        hr = cfg.header_row
        headers = {c: _clean(ws.cell(hr, c).value) for c in range(1, ws.max_column + 1)}
        languages: list[Language] = []
        pending_revisions: list[tuple[int, str]] = []
        for col, header in headers.items():
            if not header:
                continue
            base = _lookup(cfg.revisions, header)
            if base is not None or _REVISED_RE.search(header):
                pending_revisions.append((col, header))
                continue
            code = _lookup(cfg.languages, header) or DEFAULT_LANG_CODES.get(header.strip().lower())
            if code is None:
                code = re.sub(r"\W+", "_", header.strip().lower()).strip("_")
            languages.append(Language(code=code, name=header.strip(), column=col))

        # 修訂欄：設定檔指定對應語言，否則視為左邊最近的語言欄的修訂版
        for col, header in pending_revisions:
            base_name = _lookup(cfg.revisions, header)
            target = None
            if base_name is not None:
                target = next((l for l in languages
                               if l.name.lower() == base_name.strip().lower() or l.code == base_name), None)
            if target is None:
                left = [l for l in languages if l.column < col]
                target = left[-1] if left else None
            if target is not None:
                target.revision_column = col
                target.revision_name = header.strip()

        source = cfg.source_language
        if source not in [l.code for l in languages]:
            source = languages[0].code
        src_col = next(l.column for l in languages if l.code == source)

        entries: list[Entry] = []
        for r in range(hr + 1, ws.max_row + 1):
            row_values = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
            if not any(_clean(v) for v in row_values):
                continue
            key = _clean(ws.cell(r, src_col).value) or f"(第 {r} 列)"
            entry = Entry(row=r, key=key)
            for lang in languages:
                text = _clean(ws.cell(r, lang.column).value)
                cell = f"{get_column_letter(lang.column)}{r}"
                tr = Translation(lang=lang.code, text=text, cell=cell, column_name=lang.name)
                if lang.revision_column:
                    revised = _clean(ws.cell(r, lang.revision_column).value)
                    if revised:
                        tr = Translation(
                            lang=lang.code, text=revised,
                            cell=f"{get_column_letter(lang.revision_column)}{r}",
                            column_name=lang.revision_name or lang.name,
                            old_text=text, old_cell=cell,
                        )
                entry.translations[lang.code] = tr
            entries.append(entry)

        return TranslationSheet(path=Path(cfg.excel), sheet=ws.title, source_lang=source,
                                languages=languages, entries=entries)
    finally:
        wb.close()
