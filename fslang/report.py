"""輸出 Excel 報表。

工作表：
- 摘要      每個語言各狀態的數量與正確率
- 對照表    和原本翻譯 Excel 一樣的版面，儲存格依結果上色，註解寫遊戲實際文字
- 明細      每一句、每個語言一列，差異處以紅字標示
- Excel檢查 翻譯檔本身的問題（lint）
"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .excel import TranslationSheet
from .lint import ERROR, Issue
from .matcher import STATUS_LABELS, STATUS_ORDER, Result, summarize

STATUS_FILLS = {
    "PASS": "C6EFCE",
    "WARN": "FFEB9C",
    "OUTDATED": "FFC7CE",
    "WRONG_LANG": "FFC7CE",
    "MISMATCH": "FFC7CE",
    "MISSING": "F8CBAD",
    "NO_CAPTURE": "D9D9D9",
    "NO_EXPECTED": "F2F2F2",
}
SEVERITY_FILLS = {ERROR: "FFC7CE", "WARN": "FFEB9C"}

_HEADER_FILL = PatternFill("solid", fgColor="FFFF00")
_HEADER_FONT = Font(bold=True)
_WRAP = Alignment(wrap_text=True, vertical="top")
_THIN = Side(style="thin", color="BFBFBF")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_RED = InlineFont(color="FFC00000", b=True)


def _clean(value):
    if isinstance(value, str):
        return ILLEGAL_CHARACTERS_RE.sub("", value)
    return value


def _fill(color: str) -> PatternFill:
    return PatternFill("solid", fgColor=color)


def _header(ws, row: int, values: list[str], widths: list[int] | None = None):
    for i, v in enumerate(values, start=1):
        c = ws.cell(row, i, v)
        c.fill, c.font, c.alignment, c.border = _HEADER_FILL, _HEADER_FONT, _WRAP, _BORDER
        if widths:
            ws.column_dimensions[get_column_letter(i)].width = widths[i - 1]


def _rich(segments: list[tuple[str, str]]) -> CellRichText | str:
    """segments: [(text, highlight?)]"""
    parts = []
    for text, hl in segments:
        text = _clean(text)
        if not text:
            continue
        parts.append(TextBlock(_RED, text) if hl else text)
    if not parts:
        return ""
    if all(isinstance(p, str) for p in parts):
        return "".join(parts)
    return CellRichText(parts)


def _diff_cells(segs: list[tuple[str, str, str]]):
    hl = ("replace", "delete", "insert")
    exp = [(a, tag in hl) for tag, a, b in segs if a]
    act = [(b, tag in hl) for tag, a, b in segs if b]
    return _rich(exp), _rich(act)


def _write_summary(ws, sheet: TranslationSheet, results: list[Result], issues: list[Issue] | None,
                   meta: dict):
    ws["A1"] = f"多語系比對報告 - {meta.get('game_name', sheet.sheet)}"
    ws["A1"].font = Font(bold=True, size=14)
    info = [
        ("Excel", str(sheet.path.name)),
        ("工作表", sheet.sheet),
        ("擷取資料", meta.get("captures", "")),
        ("產生時間", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
    ]
    for i, (k, v) in enumerate(info, start=2):
        ws.cell(i, 1, k).font = _HEADER_FONT
        ws.cell(i, 2, _clean(v))
    start = len(info) + 3
    cols = ["語言", "代碼"] + [STATUS_LABELS[s] for s in STATUS_ORDER] + ["正確率"]
    _header(ws, start, cols, [16, 8] + [12] * len(STATUS_ORDER) + [10])
    summary = summarize(results)
    r = start
    for lang in sheet.languages:
        if lang.code not in summary:
            continue
        r += 1
        counts = summary[lang.code]
        ws.cell(r, 1, lang.name)
        ws.cell(r, 2, lang.code)
        for j, status in enumerate(STATUS_ORDER, start=3):
            c = ws.cell(r, j, counts[status])
            if counts[status] and status not in ("PASS", "NO_EXPECTED"):
                c.fill = _fill(STATUS_FILLS[status])
        checked = sum(counts[s] for s in STATUS_ORDER if s not in ("NO_CAPTURE", "NO_EXPECTED"))
        rate = ws.cell(r, len(cols), (counts["PASS"] + counts["WARN"]) / checked if checked else None)
        rate.number_format = "0.0%"
    for row in ws.iter_rows(min_row=start, max_row=r, max_col=len(cols)):
        for c in row:
            c.border = _BORDER

    r += 2
    ws.cell(r, 1, "顏色說明").font = _HEADER_FONT
    notes = {
        "PASS": "遊戲文字和 Excel 完全相同",
        "WARN": "內容相同，但大小寫/空白/全半形/引號樣式不同",
        "OUTDATED": "遊戲還在用被 (revised) 欄取代的舊翻譯",
        "WRONG_LANG": "遊戲顯示的是其他語言（例如還是英文）",
        "MISMATCH": "找到相似的文字但內容不同，差異見「明細」紅字（說明欄格式：「Excel」→「遊戲」）",
        "MISSING": "遊戲文字中找不到這句",
        "NO_CAPTURE": "這個語言沒有擷取資料",
        "NO_EXPECTED": "Excel 這格沒有翻譯",
    }
    for status in STATUS_ORDER:
        r += 1
        c = ws.cell(r, 1, STATUS_LABELS[status])
        c.fill = _fill(STATUS_FILLS[status])
        ws.cell(r, 2, notes[status])
    if issues is not None:
        r += 2
        errors = sum(1 for i in issues if i.severity == ERROR)
        ws.cell(r, 1, "Excel 檢查").font = _HEADER_FONT
        ws.cell(r, 2, f"錯誤 {errors} 筆、警告 {len(issues) - errors} 筆（見「Excel檢查」工作表）")


def _write_matrix(ws, sheet: TranslationSheet, results: list[Result]):
    langs = [l for l in sheet.languages if any(r.lang == l.code for r in results)]
    _header(ws, 1, ["Excel列", sheet.language(sheet.source_lang).name + "（原文）"] + [l.name for l in langs],
            [8, 40] + [36] * len(langs))
    by_key = {(r.row, r.lang): r for r in results}
    for i, entry in enumerate(sheet.entries, start=2):
        ws.cell(i, 1, entry.row).alignment = _WRAP
        ws.cell(i, 2, _clean(entry.key)).alignment = _WRAP
        for j, lang in enumerate(langs, start=3):
            res = by_key.get((entry.row, lang.code))
            if res is None:
                continue
            c = ws.cell(i, j, _clean(res.expected) or "")
            c.alignment, c.border = _WRAP, _BORDER
            c.fill = _fill(STATUS_FILLS[res.status])
            if res.status != "PASS":
                lines = [f"【{res.label}】{res.cell}"]
                if res.actual:
                    lines.append(f"遊戲：{res.actual}")
                if res.note:
                    lines.append(f"說明：{res.note}")
                text = _clean("\n".join(lines))[:2000]
                comment = Comment(text, "fslang")
                comment.width, comment.height = 360, 160
                c.comment = comment
    ws.freeze_panes = "C2"


def _write_details(ws, sheet: TranslationSheet, results: list[Result]):
    cols = ["Excel列", "儲存格", "語言", "狀態", "原文(EN)", "Excel 標準翻譯", "遊戲實際文字", "相似度", "說明", "來源"]
    _header(ws, 1, cols, [8, 9, 12, 14, 40, 48, 48, 9, 48, 40])
    for i, r in enumerate(results, start=2):
        ws.cell(i, 1, r.row)
        ws.cell(i, 2, r.cell)
        ws.cell(i, 3, r.lang_name)
        st = ws.cell(i, 4, r.label)
        st.fill = _fill(STATUS_FILLS[r.status])
        ws.cell(i, 5, _clean(r.key))
        if r.diff:
            exp, act = _diff_cells(r.diff)
            ws.cell(i, 6).value = exp
            ws.cell(i, 7).value = act
        else:
            ws.cell(i, 6, _clean(r.expected))
            ws.cell(i, 7, _clean(r.actual or ""))
        sim = ws.cell(i, 8, r.similarity)
        sim.number_format = "0%"
        ws.cell(i, 9, _clean(r.note))
        ws.cell(i, 10, _clean(r.source or ""))
        for j in range(1, len(cols) + 1):
            ws.cell(i, j).alignment = _WRAP
            ws.cell(i, j).border = _BORDER
    ws.freeze_panes = "E2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{max(1, len(results) + 1)}"


def _write_lint(ws, issues: list[Issue]):
    cols = ["Excel列", "儲存格", "語言", "等級", "規則", "說明", "翻譯文字", "原文(EN)"]
    _header(ws, 1, cols, [8, 9, 12, 8, 13, 50, 60, 50])
    for i, it in enumerate(issues, start=2):
        values = [it.row, it.cell, it.lang_name, it.severity, it.rule, it.message, it.text, it.key]
        for j, v in enumerate(values, start=1):
            c = ws.cell(i, j, _clean(v))
            c.alignment, c.border = _WRAP, _BORDER
        ws.cell(i, 4).fill = _fill(SEVERITY_FILLS.get(it.severity, "FFFFFF"))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{max(1, len(issues) + 1)}"


def write_compare_report(path: Path, sheet: TranslationSheet, results: list[Result],
                         issues: list[Issue] | None = None, meta: dict | None = None) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    _write_summary(wb.active, sheet, results, issues, meta or {})
    wb.active.title = "摘要"
    _write_matrix(wb.create_sheet("對照表"), sheet, results)
    _write_details(wb.create_sheet("明細"), sheet, results)
    if issues is not None:
        _write_lint(wb.create_sheet("Excel檢查"), issues)
    wb.save(path)
    return path


def write_lint_report(path: Path, sheet: TranslationSheet, issues: list[Issue]) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Excel檢查"
    _write_lint(ws, issues)
    wb.save(path)
    return path
