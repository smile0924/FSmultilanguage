"""端對端測試：用 Playwright 打開本機假遊戲（examples/demo_game）擷取文字，再和 Excel 比對。"""

import functools
import threading
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path

import openpyxl
import pytest

from fslang.collector import collect
from fslang.config import load_config
from fslang.corpus import load_captures
from fslang.excel import load_sheet
from fslang.matcher import MISMATCH, MISSING, OUTDATED, PASS, WARN, WRONG_LANG, Matcher, compare
from fslang.report import write_compare_report

ROOT = Path(__file__).resolve().parents[1]
pytest.importorskip("playwright")


class _QuietHandler(SimpleHTTPRequestHandler):
    def log_message(self, *args):
        pass


@pytest.fixture(scope="module")
def demo_url():
    handler = functools.partial(_QuietHandler, directory=str(ROOT / "examples" / "demo_game"))
    server = ThreadingHTTPServer(("127.0.0.1", 0), handler)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    yield f"http://127.0.0.1:{server.server_address[1]}/index.html?lang={{game_lang}}"
    server.shutdown()


def test_collect_and_compare_demo(demo_url, tmp_path):
    cfg = load_config(ROOT / "configs" / "demo_tree_of_wild.json")
    cfg.collect["game_url"] = demo_url
    sheet = load_sheet(cfg)
    langs = ["en", "th", "ja", "ko", "it", "es", "fr", "bn"]
    try:
        saved = collect(cfg, sheet, langs, tmp_path / "cap")
    except Exception as e:  # 沒有瀏覽器的環境
        pytest.skip(f"無法啟動瀏覽器：{e}")
    assert len(saved) == len(langs)

    corpora = load_captures(tmp_path / "cap", langs)
    m = cfg.match
    results = compare(sheet, corpora, Matcher(cfg.placeholders, m["wildcard_max_chars"], m["similarity_threshold"]),
                      langs)
    by = {(r.lang, r.row): r for r in results}

    # 擷取到的是畫面上的說明文字（DOM），不是只有語系檔
    assert by[("en", 2)].source.startswith("DOM")
    # 正確的語言
    assert {r.status for r in results if r.lang == "en"} == {PASS}
    # fr：NBSP 換成一般空白仍算正確
    assert {r.status for r in results if r.lang == "fr"} == {PASS}
    # 故意放的錯誤
    assert by[("th", 3)].status == OUTDATED
    assert by[("ja", 4)].status == MISMATCH
    assert by[("ko", 18)].status == WRONG_LANG
    assert by[("it", 9)].status == MISSING
    assert by[("es", 2)].status == WARN
    # bn 沒有語系檔，遊戲退回英文
    assert {r.status for r in results if r.lang == "bn"} == {WRONG_LANG}
    # 其他句子都正確
    expected_bad = {("th", 3), ("ja", 4), ("ko", 18), ("it", 9), ("es", 2)}
    assert all(r.status == PASS for r in results if r.lang != "bn" and (r.lang, r.row) not in expected_bad)

    out = write_compare_report(tmp_path / "report.xlsx", sheet, results)
    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames == ["摘要", "對照表", "明細"]
