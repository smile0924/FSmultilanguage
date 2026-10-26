import openpyxl
import pytest

from fslang.config import config_from_args
from fslang.corpus import Corpus, Record
from fslang.excel import load_sheet
from fslang.matcher import (MISMATCH, MISSING, NO_CAPTURE, NO_EXPECTED, OUTDATED, PASS, WARN, WRONG_LANG, Matcher,
                            compare, diff_segments, diff_summary)

PLACEHOLDERS = [r"\[[^\]]*\]", r"(?<![A-Za-z])X(?![A-Za-z])"]


@pytest.fixture()
def sheet(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Game"
    ws.append(["EN", "Thai", "TH (revised)", "Japan", "French"])
    ws.append(["Unlock Free Game", "ปลดล็อกเกมฟรี", None, "フリーゲームをアンロック", "Débloquez le jeu gratuit"])
    ws.append(["Get Free Spins!", "รับหมุนฟรี!", "รับฟรีสปิน!", "フリースピンを獲得！", "Obtenez des tours gratuits !"])
    ws.append(["[WILD] climbs up every winning", "[WILD] ปีนขึ้นทุกครั้งที่ได้รางวัล", None,
               "[WILD]はすべての勝利を上げます", "Le [WILD] monte à chaque gain"])
    ws.append(["Win Multiplier up to X in the Free Game", "ตัวคูณสูงสุด X ในเกมฟรี", None,
               "フリーゲームでマルチプライヤーを最大Xまで獲得", "Multiplicateur allant jusqu'à X"])
    ws.append(["Tier 1", "ระดับ 1", None, "ティア1", None])
    path = tmp_path / "t.xlsx"
    wb.save(path)
    return load_sheet(config_from_args(str(path), "Game"))


def run(sheet, lang, *texts):
    corpus = Corpus([Record("test", t) for t in texts])
    results = compare(sheet, {lang: corpus}, Matcher(PLACEHOLDERS), [lang])
    return {r.row: r for r in results}


def test_all_correct_with_placeholders_and_line_breaks(sheet):
    res = run(sheet, "ja",
              "メニュー\nフリーゲームをアンロック",
              "フリースピンを獲得！",
              "はすべての勝利を上げます",                     # [WILD] 是圖示，畫面上沒有字
              "フリーゲームでマルチプライヤーを最大x2,000まで獲得",  # X 是實際倍數
              "ティア1")
    assert {r.status for r in res.values()} == {PASS}


def test_sentence_split_across_lines_still_matches(sheet):
    res = run(sheet, "fr", "Obtenez des tours\ngratuits\u00a0!")
    assert res[3].status == PASS


def test_warn_on_case_difference(sheet):
    res = run(sheet, "fr", "DÉBLOQUEZ LE JEU GRATUIT")
    assert res[2].status == WARN
    assert res[2].actual == "DÉBLOQUEZ LE JEU GRATUIT"
    assert "大小寫" in res[2].note


def test_outdated_translation_detected(sheet):
    res = run(sheet, "th", "รับหมุนฟรี!")
    assert res[3].status == OUTDATED
    assert res[3].cell == "C3" or "C3" in res[3].note


def test_wrong_language_detected(sheet):
    res = run(sheet, "th", "Tier 1")
    assert res[6].status == WRONG_LANG
    assert "EN" in res[6].note


def test_mismatch_reports_diff(sheet):
    res = run(sheet, "ja", "はすべで勝利を上げます", "フリースピンを獲得！")
    assert res[4].status == MISMATCH
    assert res[4].note == "「ての」→「で」"
    assert res[3].status == PASS


def test_missing_and_no_expected_and_no_capture(sheet):
    res = run(sheet, "fr", "Something unrelated")
    assert res[2].status == MISSING
    assert res[6].status == NO_EXPECTED
    results = compare(sheet, {}, Matcher(PLACEHOLDERS), ["ja"])
    assert {r.status for r in results} == {NO_CAPTURE}


def test_fuzzy_does_not_pick_other_rows_exact_text(sheet):
    # 第 2 列不在遊戲裡；畫面上只有第 3 列的文字，不應拿第 3 列來當第 2 列的「內容不符」
    res = run(sheet, "en", "Get Free Spins!")
    assert res[3].status == PASS
    assert res[2].status == MISSING


def test_diff_ignores_placeholders():
    import re
    segs = diff_segments("Win up to X now", "Win up to x500 now", [re.compile(r"(?<![A-Za-z])X(?![A-Za-z])")])
    assert diff_summary(segs) == ""
    segs = diff_segments("[WILD]はすべての", "はすべで", [re.compile(r"\[[^\]]*\]")])
    assert diff_summary(segs) == "「ての」→「で」"
